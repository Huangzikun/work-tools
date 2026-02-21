"""
创建测试固件文件
"""
import sys
from pathlib import Path

# 添加项目根目录到路径
project_root = Path(__file__).parent.parent.parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_simple_template():
    """创建简单模板"""
    doc = Document()

    # 添加标题
    title = doc.add_heading('学生信息表', 0)

    # 添加段落（带格式）
    p = doc.add_paragraph()
    run = p.add_run('学生姓名：{{姓名}}')
    run.font.size = Pt(12)
    run.font.name = '宋体'
    run.bold = True

    p = doc.add_paragraph()
    run = p.add_run('学号：{{学号}}')
    run.font.size = Pt(12)
    run.font.name = '宋体'

    p = doc.add_paragraph()
    run = p.add_run('课程：{{课程}}')
    run.font.size = Pt(12)
    run.font.name = '宋体'

    p = doc.add_paragraph()
    run = p.add_run('课时：{{课时}}')
    run.font.size = Pt(12)
    run.font.name = '宋体'

    # 保存
    output_path = Path(__file__).parent / 'fixtures' / 'template_simple.docx'
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"已创建: {output_path}")


def create_table_template():
    """创建表格模板"""
    doc = Document()

    # 添加标题
    title = doc.add_heading('课程信息表', 0)

    # 添加表格
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'

    # 填充表格
    cells_data = [
        ('课程名称', '{{课程}}'),
        ('授课教师', '{{教师}}'),
        ('课时数', '{{课时}}'),
        ('学分', '{{学分}}'),
        ('课程类型', '{{类型}}'),
    ]

    for i, (label, placeholder) in enumerate(cells_data):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = placeholder

    # 保存
    output_path = Path(__file__).parent / 'fixtures' / 'template_table.docx'
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"已创建: {output_path}")


def create_multiple_template():
    """创建多次出现占位符的模板（用于测试全量/首次替换）"""
    doc = Document()

    # 添加标题
    title = doc.add_heading('学生信息表（重复测试）', 0)

    # 第一次出现
    p = doc.add_paragraph()
    run = p.add_run('学生姓名：{{姓名}}')
    run.font.size = Pt(12)
    run.bold = True

    p = doc.add_paragraph()
    run = p.add_run('学号：{{学号}}')
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    run = p.add_run('课程：{{课程}}')
    run.font.size = Pt(12)

    # 第二次出现（同名占位符）
    p = doc.add_paragraph()
    run = p.add_run('授课教师：{{姓名}}')
    run.font.size = Pt(12)
    run.italic = True

    p = doc.add_paragraph()
    run = p.add_run('班主任：{{姓名}}')
    run.font.size = Pt(12)
    run.underline = True

    # 保存
    output_path = Path(__file__).parent / 'fixtures' / 'template_multiple.docx'
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"已创建: {output_path}")


if __name__ == '__main__':
    print("开始创建测试固件文件...")
    create_simple_template()
    create_table_template()
    create_multiple_template()
    print("完成！")
