import os
from docx import Document
from pathlib import Path
import pandas as pd
from common.llm_client import LLMClient


def extract_table_text(doc):

    full_text = []

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    full_text.append(paragraph.text)

    return "\n".join(list(dict.fromkeys(full_text)))

# 尝试读取文件内容
document = Document("/Users/huangzikun/Desktop/桂林学院/算法设计与分析/期末材料/计科/2022计算机科学与技术《算法设计与分析》实验报告黄子坤55份/202213008201计算机科学与技术张耀匀/202213008201张耀匀-实验实训报告1_1.docx")
document_text = extract_table_text(document)

teacher = "实验报告针对归并与快速排序实验，目的应包含分治算法的基本思想；实验原理应包括分治、合并和解决；结果分析应对比两种排序。若每项内容均包含并符合算法思想，应得"

client = LLMClient()

print("----- standard request -----")
output_text = client.generate(
    system_prompt="你是桂林学院信息工程学院的一名计算机专任教师，你拥有丰富的教学经验。你的任务是针对学生提交的实验报告进行批改。你应该理解、使用用户提交的“教师要求”部分对“学生作答”部分进行批阅。你可以选择的分数为70,80,90和100分，并给出一个50字以内的批阅评语。生成json格式的内容，包含一个score和一个comment字段。",
    user_prompt=f"教师要求:{teacher};学生作答:{document_text}",
    json_output=True,
)

print(output_text)
