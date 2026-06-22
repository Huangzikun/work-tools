from docx import Document
from common.llm_client import LLMClient

def read_docx(file_path):
    """读取docx文件内容并返回文本"""
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    return '\n'.join(full_text)


client = LLMClient()

type = "{' 课次 ',' 授课内容 ',' 授课学时 ',' 知识目标 ',' 能力目标 ',' 情感与价值观目标 ',' 重点 ',' 难点 ',' 教学工具或资源 ',' 课堂导入 ',' 课堂导入时间分配 ',' 课堂内容 ',' 课堂内容时间分配 ',' 教学方法与设计 ',' 课堂小结 ',' 课堂小结时间分配 ',' 课后作业 ',' 课后作业时间分配 ',' 教学反思 '}"

doc_content = read_docx("/Users/huangzikun/PycharmProjects/work-tools/第1章 Java开发入门.docx")
print("----- standard request -----")
output_text = client.generate(
    system_prompt="你是桂林学院信息工程学院的一名计算机专任教师，你拥有丰富的教学经验。",
    user_prompt=f"从文件中读取以下列表中的每一项，生成json格式的内容。列表内容如下：\n{type}\n。文档内容如下：\n{doc_content}\n如果文档未提及，则json内不包含这项内容。如果一项中包含一个列表，则每一个条目后面增加换行。",
)

print(output_text)
