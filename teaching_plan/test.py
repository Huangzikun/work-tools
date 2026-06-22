from docx import Document
import json
import os
from common.llm_client import LLMClient

def read_docx(file_path):
    """读取docx文件内容并返回文本"""
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    return '\n'.join(full_text)


client = LLMClient()

type = "{'课次,应为一个整数，如1','授课内容,应精简描述知识点','授课学时,应为一个整数，如3','知识目标','能力目标','情感与价值观目标','重点','难点','教学工具或资源','课堂导入','课堂导入时间分配','课堂内容，应尽可能的详细、准确描述授课内容和知识点','课堂内容时间分配','教学方法与设计','课堂小结','课堂小结时间分配','课后作业','课后作业时间分配','教学反思'}"


file_dir = "/Users/huangzikun/Desktop/test/"

file_list = os.listdir(file_dir)
file_list.sort()

res_list = []
for file in file_list:
    doc_content = read_docx(os.path.join(file_dir, file))

    print("----- standard request -----")
    output_text = client.generate(
        system_prompt="你是桂林学院信息工程学院的一名计算机专任教师，你拥有丰富的教学经验。你的任务是按照新的格式生成新的教案列表。从文件中读取以下列表中的每一项，生成json格式的内容。json应为一个数组，其中每一个对象包括后续列表内的每一项。",
        user_prompt=f"根据文档内容中的教学时数/3确定教案的数量，如教学时数为9，即你应该给我提供9/3=3个教案的内容，最多不超过4个教案。列表内容如下：\n{type}\n。文档内容如下：\n{doc_content}\n如果文档未提及，则对象内的这一项设置为空字符串。如果一项中包含一个列表，则每一个条目后面增加换行。",
        json_output=True,
    )

    print(output_text)
    res_list.extend(json.loads(output_text))

with open("temp.txt", "w", encoding="utf-8") as f:
    json.dump(res_list, f, ensure_ascii=False, indent=2)
