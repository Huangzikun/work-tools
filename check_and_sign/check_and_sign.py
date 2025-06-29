import json
from volcenginesdkarkruntime import Ark
import os
from docx import Document
from docx.shared import Cm
import argparse



def extract_table_text(doc):

    full_text = []

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    full_text.append(paragraph.text)

    return "\n".join(list(dict.fromkeys(full_text)))

def sign_by_picture(file_path, save_path, score, review):
    doc = Document(file_path)
    if len(doc.tables) <= 0:
        print(f"can't sign file. file = {file_path}")
        return

    for table in doc.tables:
        for row in table.rows:
            for i in range(len(row.cells)):

                if str.strip(row.cells[i].text) == '教师评阅':

                    row.cells[i + 1].text = (f"{review}"
                                             f"成绩：{score}分"
                                             f"{os.linesep}"
                                             f"{os.linesep}"
                                             f"{os.linesep}"
                                             f"{os.linesep}"
                                             f"{os.linesep}"
                                             f"                       教师签名：{sign}   {sign_date}")
                    row.cells[i + 1].add_paragraph().add_run().add_picture(sign_picture, width=Cm(2))
                    doc.save(save_path)
                    print(f"sign successfully. file = {file_path}")

                    return

    print(f"sign fail. file = {file_path}")

parser = argparse.ArgumentParser()
parser.add_argument('--sign_picture', type=str, help='签名图片地址', required=True)
parser.add_argument('--sign', type=str, help='签名字符串，如某某某', required=True)
parser.add_argument('--sign_date', type=str, help='签名时间', required=True)

args = parser.parse_args()
sign_picture = args.sign_picture
sign = args.sign
sign_date = args.sign_date

# 尝试读取文件内容
file_path = "/Users/huangzikun/Desktop/桂林学院/算法设计与分析/期末材料/计科/2022计算机科学与技术《算法设计与分析》实验报告黄子坤55份/202213008201计算机科学与技术张耀匀/202213008201张耀匀-实验实训报告1_1.docx"
document = Document(file_path)
document_text = extract_table_text(document)

teacher = "实验报告针对归并与快速排序实验，目的应包含分治算法的基本思想；实验原理应包括分治、合并和解决；结果分析应对比两种排序。若每项内容均包含并符合算法思想，应得"

# 请确保您已将 API Key 存储在环境变量 ARK_API_KEY 中
# 初始化Ark客户端，从环境变量中读取您的API Key
client = Ark(
    # 此为默认路径，您可根据业务所在地域进行配置
    base_url="https://ark.cn-beijing.volces.com/api/v3",
    # 从环境变量中获取您的 API Key。此为默认方式，您可根据需要进行修改
    api_key=os.environ.get("ARK_API_KEY"),
)

# Non-streaming:
print("----- standard request -----")
completion = client.chat.completions.create(
   # 指定您创建的方舟推理接入点 ID，此处已帮您修改为您的推理接入点 ID
    model="doubao-seed-1-6-250615",
    messages=[
        {"role": "system", "content": "你是桂林学院信息工程学院的一名计算机专任教师，你拥有丰富的教学经验。你的任务是针对学生提交的实验报告进行批改。你应该理解、使用用户提交的“教师要求”部分对“学生作答”部分进行批阅。你可以选择的分数为70,80,90和100分，并给出一个50字以内的批阅评语。生成json格式的内容，包含一个score和一个comment字段。"},
        {"role": "user", "content": f"教师要求:{teacher};学生作答:{document_text}"},
    ],
    response_format={
        "type": "json_object",
    }
)

print(completion.choices[0].message.content)
result = json.loads(completion.choices[0].message.content)
sign_by_picture(file_path, file_path, result["score"], result["comment"])
