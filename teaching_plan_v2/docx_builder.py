"""
Word文档构建器 - 将所有课时合并到一个教案文档中
"""
import os
import shutil
import logging
from typing import List
from pathlib import Path
from docx import Document

from common.docx_template import DocxTemplateReplacer, ReplacerConfig
from .models import LessonPlan
from .text_formatter import format_all_lesson_data

logger = logging.getLogger(__name__)


class DocxBuilder:
    """Word文档构建器"""

    def __init__(self, template_path: str, output_dir: str):
        """
        初始化文档构建器

        Args:
            template_path: Word模板文件路径
            output_dir: 输出目录
        """
        self.template_path = Path(template_path)
        self.output_dir = Path(output_dir)

        # 确保输出目录存在
        self.output_dir.mkdir(parents=True, exist_ok=True)

        logger.info(f"文档构建器初始化完成，模板: {template_path}")

    def build_all(
        self,
        lessons: List[LessonPlan],
        output_filename: str = "教案.docx"
    ) -> str:
        """
        将所有课时合并到一个教案文档中

        Args:
            lessons: 所有课时列表
            output_filename: 输出文件名

        Returns:
            输出文件路径
        """
        if not lessons:
            raise ValueError("课时列表不能为空")

        logger.info(f"开始构建合并教案文档，共{len(lessons)}个课时")

        output_path = self.output_dir / output_filename

        # 复制模板作为基础文档
        temp_template = self.output_dir / "_temp_template.docx"
        shutil.copy(self.template_path, temp_template)

        # 首次替换模式：每个课时只替换第一次出现的占位符
        config = ReplacerConfig(replace_all=False)

        try:
            current_doc = str(temp_template)

            for i, lesson in enumerate(lessons):
                logger.info(f"处理第{i+1}/{len(lessons)}个课时: {lesson}")

                # 为每个课时创建替换器
                replacer = DocxTemplateReplacer(current_doc, config)

                # 输出文件（每次使用新文件名避免锁定）
                if i < len(lessons) - 1:
                    # 还有下一个课时，使用临时文件
                    next_doc = self.output_dir / f"_temp_step_{i}.docx"
                else:
                    # 最后一个课时，直接使用最终输出路径
                    next_doc = output_path

                # 格式化教案数据（在编号步骤之间添加换行）
                formatted_data = format_all_lesson_data(lesson.to_dict())

                # 替换当前课时的内容
                replacer.replace(formatted_data, str(next_doc))

                # 获取统计信息
                stats = replacer.get_stats()
                logger.debug(f"  替换统计: 找到{stats['placeholders_found']}个占位符，替换{stats['placeholders_replaced']}个")

                # 更新当前文档路径，删除旧的临时文件
                if i > 0 and current_doc != str(temp_template):
                    try:
                        Path(current_doc).unlink()
                    except Exception:
                        pass
                current_doc = str(next_doc)

            # 清理文档：删除未使用的表格和空白节
            self._cleanup_document(str(output_path), len(lessons))

            logger.info(f"教案文档生成完成: {output_path}")
            return str(output_path)

        except Exception as e:
            logger.error(f"生成教案文档失败: {e}")
            # 清理临时文件
            if temp_template.exists():
                temp_template.unlink()
            # 清理步骤临时文件
            for i in range(len(lessons)):
                temp_file = self.output_dir / f"_temp_step_{i}.docx"
                if temp_file.exists():
                    temp_file.unlink()
            raise

    def _cleanup_document(self, doc_path: str, used_lesson_count: int):
        """
        清理文档：删除未使用的表格、空白段落和多余节

        Args:
            doc_path: 文档路径
            used_lesson_count: 实际使用的课时数
        """
        doc = Document(doc_path)

        # 1. 删除未使用的课时表格（保留前2个基本信息表格 + used_lesson_count个课时表格）
        # 表格0-1是课程基本信息，表格2开始是课时表格
        total_tables = len(doc.tables)
        target_table_count = 2 + used_lesson_count  # 保留的表格数

        if total_tables > target_table_count:
            # 从后往前删除未使用的表格
            for i in range(total_tables - 1, target_table_count - 1, -1):
                try:
                    table = doc.tables[i]
                    table._element.getparent().remove(table._element)
                    logger.debug(f"删除未使用的表格{i}")
                except Exception as e:
                    logger.debug(f"删除表格{i}时出错: {e}")

        # 2. 删除末尾的空白段落
        para_count = len(doc.paragraphs)
        for i in range(para_count - 1, -1, -1):
            para = doc.paragraphs[i]
            if not para.text.strip():
                try:
                    p_element = para._element
                    if p_element.getparent() is not None:
                        p_element.getparent().remove(p_element)
                except Exception:
                    continue
            else:
                break

        # 3. 删除多余的节（保留前2个节）
        sections = doc.sections
        while len(sections) > 2:
            try:
                sections[-1]._element.getparent().remove(sections[-1]._element)
                sections = doc.sections
            except Exception:
                break

        doc.save(doc_path)
        logger.info(f"文档清理完成: {len(doc.tables)}个表格, {len(doc.sections)}个节, {len(doc.paragraphs)}个段落")
