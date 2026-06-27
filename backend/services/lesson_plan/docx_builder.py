"""Word 文档构建器（参数化版，复制自 teaching_plan_v2/docx_builder.py 并加首页表格填充）。

改造点：
1. 构造函数只接收 template_path（不再管理 output_dir，输出由调用方指定）
2. 新增 fill_first_page_tables(doc, course_info, teacher_info) — 按表格 cell 索引覆写首页静态字段
3. build_all(lessons, output_path, course_info, teacher_info) — 接收表单数据，输出路径直接由调用方指定
"""

from __future__ import annotations

import logging
import os
import shutil
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from common.docx_template import DocxTemplateReplacer, ReplacerConfig

from .models import LessonPlan
from .text_formatter import format_all_lesson_data

logger = logging.getLogger(__name__)


def _set_cell_text(cell, text: str) -> None:
    """覆写 cell 内容，保留第一段格式（清空 runs 后 add_run）。"""
    text = "" if text is None else str(text)
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        cell.text = text
        return
    first = paragraphs[0]
    for run in list(first.runs):
        run._element.getparent().remove(run._element)
    first.add_run(text)
    for p in paragraphs[1:]:
        p._element.getparent().remove(p._element)


def fill_first_page_tables(doc: Document, course_info: dict, teacher_info: dict) -> None:
    """填充首页表格 0（课程基本信息）和表格 1（教师信息）。

    表格 0 布局（4 行 8 列，部分合并）：
      r0c0=课程名称(label)  r0c1=值(合并 c1-3)  r0c4=英文名称(label)  r0c5=值(合并 c5-7)
      r1c0=学分   r1c1=值   r1c2=理论学时  r1c3=值   r1c4=实践学时  r1c5=值   r1c6=线上学时  r1c7=值
      r2c0=适用专业  r2c1=值(合并)  r2c4=先修课程  r2c5=值(合并)
      r3c0=开课学期  r3c1=值(合并)  r3c4=课程教研室  r3c5=值(合并)
    表格 1 布局（5 行 2 列）：
      r0=授课教师  r1=所属单位  r2=撰写日期  r3=课程类型  r4=课程性质
    """
    tables = doc.tables
    if len(tables) < 2:
        logger.warning("首页表格数量不足: %d", len(tables))
        return

    t0 = tables[0]
    rows0 = t0.rows
    if len(rows0) >= 1:
        _set_cell_text(rows0[0].cells[1], course_info.get("课程名称", ""))
        _set_cell_text(rows0[0].cells[5], course_info.get("英文名称", ""))
    if len(rows0) >= 2:
        _set_cell_text(rows0[1].cells[1], course_info.get("学分", ""))
        _set_cell_text(rows0[1].cells[3], course_info.get("理论学时", ""))
        _set_cell_text(rows0[1].cells[5], course_info.get("实践学时", ""))
        _set_cell_text(rows0[1].cells[7], course_info.get("线上学时", ""))
    if len(rows0) >= 3:
        _set_cell_text(rows0[2].cells[1], course_info.get("适用专业", ""))
        _set_cell_text(rows0[2].cells[5], course_info.get("先修课程", ""))
    if len(rows0) >= 4:
        _set_cell_text(rows0[3].cells[1], course_info.get("开课学期", ""))
        _set_cell_text(rows0[3].cells[5], course_info.get("课程教研室", ""))

    t1 = tables[1]
    rows1 = t1.rows
    field_map = ["授课教师", "所属单位", "撰写日期", "课程类型", "课程性质"]
    for i, field in enumerate(field_map):
        if i < len(rows1):
            _set_cell_text(rows1[i].cells[1], teacher_info.get(field, ""))


class DocxBuilder:
    def __init__(self, template_path: str):
        self.template_path = Path(template_path)
        if not self.template_path.exists():
            raise FileNotFoundError(f"模板不存在: {template_path}")
        logger.info("DocxBuilder init, template=%s", template_path)

    def build_all(
        self,
        lessons: List[LessonPlan],
        output_path: str,
        course_info: dict,
        teacher_info: dict,
    ) -> str:
        if not lessons:
            raise ValueError("课时列表不能为空")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        temp_template = output_path.parent / "_temp_template.docx"
        shutil.copy(self.template_path, temp_template)
        # 源模板可能是只读的（仓库分发），强制加写权限避免后续 save 失败
        os.chmod(temp_template, 0o644)

        # 先在临时模板上填首页表格
        doc = Document(str(temp_template))
        fill_first_page_tables(doc, course_info, teacher_info)
        doc.save(str(temp_template))

        config = ReplacerConfig(replace_all=False)

        try:
            current_doc = str(temp_template)
            for i, lesson in enumerate(lessons):
                logger.info("处理第 %d/%d 个课时: %s", i + 1, len(lessons), lesson)

                replacer = DocxTemplateReplacer(current_doc, config)

                if i < len(lessons) - 1:
                    next_doc = output_path.parent / f"_temp_step_{i}.docx"
                else:
                    next_doc = output_path

                formatted_data = format_all_lesson_data(lesson.to_dict())
                replacer.replace(formatted_data, str(next_doc))
                # 同样确保临时步骤文件可写
                os.chmod(next_doc, 0o644)

                if i > 0 and current_doc != str(temp_template):
                    try:
                        Path(current_doc).unlink()
                    except Exception:
                        pass
                current_doc = str(next_doc)

            self._cleanup_document(str(output_path), len(lessons))
            self._add_page_breaks_between_lessons(str(output_path), len(lessons))

            logger.info("教案文档生成完成: %s", output_path)
            return str(output_path)
        except Exception:
            logger.exception("生成教案失败")
            if temp_template.exists():
                temp_template.unlink(missing_ok=True)
            for i in range(len(lessons)):
                (output_path.parent / f"_temp_step_{i}.docx").unlink(missing_ok=True)
            raise

    def _cleanup_document(self, doc_path: str, used_lesson_count: int) -> None:
        doc = Document(doc_path)
        total_tables = len(doc.tables)
        target_table_count = 2 + used_lesson_count

        if total_tables > target_table_count:
            for i in range(total_tables - 1, target_table_count - 1, -1):
                try:
                    table = doc.tables[i]
                    table._element.getparent().remove(table._element)
                except Exception as e:
                    logger.debug("删除表格 %d 失败: %s", i, e)

        para_count = len(doc.paragraphs)
        for i in range(para_count - 1, -1, -1):
            para = doc.paragraphs[i]
            if not para.text.strip():
                try:
                    p_element = para._element
                    if p_element.getparent() is not None:
                        p_element.getparent().remove(p_element)
                except Exception:
                    continue
            else:
                break

        sections = doc.sections
        while len(sections) > 2:
            try:
                sections[-1]._element.getparent().remove(sections[-1]._element)
                sections = doc.sections
            except Exception:
                break

        doc.save(doc_path)

    def _add_page_breaks_between_lessons(self, doc_path: str, lesson_count: int) -> None:
        """在每个教案表格**前**插入分页符，确保：
        - 封面（含表格 0、1）与第一个教案之间分页
        - 相邻教案之间分页
        最后一个教案后不再插分页符。

        布局：封面 → 分页 → 教案1 → 分页 → 教案2 → ... → 分页 → 教案N
        """
        doc = Document(doc_path)
        first_lesson_table = 2  # 表格 0/1 是封面（课程信息 + 教师信息）
        last_lesson_table = 2 + lesson_count - 1

        try:
            # 倒序遍历：先插最后一个教案前的分页符，再插倒数第二个前的，…
            # 倒序保证已插入的分页符不影响后续表格的索引
            for table_idx in range(last_lesson_table, first_lesson_table - 1, -1):
                if table_idx >= len(doc.tables):
                    continue
                try:
                    table_element = doc.tables[table_idx]._element
                    parent = table_element.getparent()
                    table_index = parent.index(table_element)

                    p = OxmlElement("w:p")
                    pPr = OxmlElement("w:pPr")
                    p.append(pPr)
                    r = OxmlElement("w:r")
                    p.append(r)
                    br = OxmlElement("w:br")
                    br.set(qn("w:type"), "page")
                    r.append(br)
                    # 插在表格"前面"（即 table_index 位置）
                    parent.insert(table_index, p)
                except Exception as e:
                    logger.warning("在表格 %d 前添加分页符失败: %s", table_idx, e)

            doc.save(doc_path)
        except Exception as e:
            logger.warning("批量添加分页符失败: %s", e)
